# 生成「成绩查询个人码 群发」表格骨架（.xlsx，上传 Drive 时转成 Google 表格）
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROWS = 200
HEAD_FILL = PatternFill('solid', fgColor='1C1B19')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=11)
NOTE_FILL = PatternFill('solid', fgColor='FAF9F6')

wb = Workbook()

# ---------- 说明 ----------
ws = wb.active; ws.title = '说明'
lines = [
 ('成绩查询个人码 · 群发邮件', 'h1'),
 ('', ''),
 ('第 1 步：粘个人码', 'h2'),
 ('课堂系统 → 成绩 → 设置 → 学生自查 → 点「复制 学号/姓名/个人码」，', ''),
 ('然后到下面的「个人码_9E」工作表，点 A1 按 Cmd+V。10A、10H 各做一遍。', ''),
 ('', ''),
 ('第 2 步：装脚本（只做一次）', 'h2'),
 ('扩展程序 → Apps Script → 删掉里面原有内容 → 粘贴 send_codes.gs 全文 → 保存。', ''),
 ('回到表格按 Cmd+R 刷新，菜单栏会多出「成绩通知」。第一次点会让你授权，点允许。', ''),
 ('', ''),
 ('第 3 步：生成名单', 'h2'),
 ('点 成绩通知 → ① 从个人码表生成名单。班级和学号会自动填好，顺序和课堂系统一致。', ''),
 ('', ''),
 ('第 4 步：填邮箱', 'h2'),
 ('在「名单」工作表的 D 列填每个学生的邮箱。C 列中文名会自动带出——', ''),
 ('填完务必逐行看一眼：C 列的名字和 D 列的邮箱是同一个人吗？', ''),
 ('个人码发错人，那个同学就能看到别人的成绩。', ''),
 ('', ''),
 ('第 5 步：先测试，再群发', 'h2'),
 ('② 发一封给我自己 → 检查格式和网址 → ③ 发送所有还没发的。', ''),
 ('③ 可以重复点，已发送的会自动跳过；没发成功的原因写在「状态」列。', ''),
 ('', ''),
 ('注意', 'h2'),
 ('· 这张表里有学生个人码和邮箱，不要共享给别人；发完可以把个人码工作表清空。', ''),
 ('· 这个脚本只发给学生本人，不抄送家长。家长要不要收、收什么，和学校确认后再单独做。', ''),
 ('· 发之前先在课堂系统设置里勾上「对该班开放学生查询」，否则学生登录会看到「还没有开放」。', ''),
]
for i, (text, kind) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=text)
    if kind == 'h1': c.font = Font(bold=True, size=16)
    elif kind == 'h2': c.font = Font(bold=True, size=12, color='0F766E')
    else: c.font = Font(size=11)
ws.column_dimensions['A'].width = 96
ws.sheet_view.showGridLines = False

# ---------- 名单 ----------
ws = wb.create_sheet('名单')
heads = ['班级', '学号', '中文名（自动核对）', '邮箱 ← 填这一列', '状态（脚本自动填）']
widths = [10, 8, 20, 34, 30]
for j, (h, w) in enumerate(zip(heads, widths), start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.fill, c.font = HEAD_FILL, HEAD_FONT
    c.alignment = Alignment(horizontal='center')
    ws.column_dimensions[get_column_letter(j)].width = w
for r in range(2, ROWS + 2):
    ws.cell(row=r, column=3,
            value=f'=IF(OR($A{r}="",$B{r}=""),"",IFERROR(VLOOKUP($B{r},INDIRECT("个人码_"&$A{r}&"!A:C"),2,FALSE),"⚠ 没找到"))')
dv = DataValidation(type='list', formula1='"9E,10A,10H"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'A2:A{ROWS + 1}')
ws.freeze_panes = 'A2'

# ---------- 个人码 ----------
for cls in ['9E', '10A', '10H']:
    ws = wb.create_sheet(f'个人码_{cls}')
    for j, (h, w) in enumerate(zip(['学号', '姓名', '个人码'], [10, 16, 12]), start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(j)].width = w
    c = ws.cell(row=1, column=5, value=f'← 从课堂系统 {cls} 的 设置 → 学生自查 → 「复制 学号/姓名/个人码」，点 A1 粘贴（覆盖表头没关系）')
    c.font = Font(size=10, color='8A8781')
    ws.freeze_panes = 'A2'

wb.save('/private/tmp/claude-501/-Users-qingzhuling-Documents-Code/72a0c4a8-36fe-43ad-bdd2-b3a7d71cf7df/scratchpad/成绩查询个人码_群发.xlsx')
